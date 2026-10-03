"""Clean the exported USM specimen CSV while preserving uncertain source data.

This is an offline, review-oriented first pass. It does not resolve taxa against
the application database and never edits the source CSV. Unflagged records go
to cleaned.csv; records needing expert review go to needs_review.csv.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
import tempfile
import unicodedata
from collections import Counter, defaultdict
from contextlib import ExitStack
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

from openpyxl import load_workbook

REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = REPOSITORY_ROOT / "data/XLSX/Registros USM 0-300+_csv/32_310.csv"
DEFAULT_FIELDS = REPOSITORY_ROOT / "data/Campos DwC.xlsx"
DEFAULT_UBIGEO = REPOSITORY_ROOT / "backend/data/ubigeo"
DEFAULT_OUTPUT_DIR = REPOSITORY_ROOT / "data/cleaned/32_310"
DEFAULT_BATCH_OUTPUT_DIR = REPOSITORY_ROOT / "data/cleaned/usm_all"

SOURCE_COLUMNS = (
    "Código USM",
    "Tipo",
    "Familia",
    "Género",
    "Especie",
    "Autor",
    "Clasificación subespecífica",
    "Nombre vernáculo",
    "Descripción",
    "Otros",
    "País",
    "Departamento",
    "Provincia",
    "Dtto./Localidad",
    "Latitud / Northing",
    "Longitud / Easting",
    "Altitud",
    "Colector",
    "N° Colector",
    "Fecha colecta",
    "Registrador",
    "Fecha digitalización",
    "Determinado por",
    "Fecha determinacion",
    "Distribuidos en .. (Herbarios)",
    "Fecha de revisión",
)

# All priority-required terms are present in the output, including terms for
# which this source spreadsheet has no value. Those remain blank and are
# counted in summary.json rather than fabricated.
OUTPUT_DWC_FIELDS = (
    "catalogNumber",
    "typeStatus",
    "family",
    "genus",
    "specificEpithet",
    "infraspecificEpithet",
    "taxonRank",
    "scientificNameAuthorship",
    "scientificName",
    "identificationQualifier",
    "verbatimIdentification",
    "vernacularName",
    "occurrenceRemarks",
    "habitat",
    "country",
    "countryCode",
    "stateProvince",
    "county",
    "municipality",
    "locality",
    "verbatimLocality",
    "decimalLatitude",
    "decimalLongitude",
    "verbatimElevation",
    "minimumElevationInMeters",
    "maximumElevationInMeters",
    "recordedBy",
    "recordNumber",
    "verbatimEventDate",
    "eventDate",
    "year",
    "month",
    "day",
    "identifiedBy",
    "dateIdentified",
    "order",
    "originalNameUsage",
    "namePublishedIn",
    "namePublishedInYear",
    "taxonomicStatus",
    "dynamicProperties",
)
REVIEW_COLUMNS = (
    "sourceFile",
    "sourceRow",
    "catalogNumber",
    "field",
    "originalValue",
    "resultingValue",
    "severity",
    "reason",
)
AUDIT_COLUMNS = ("sourceFile", "sourceRow", "cleaningStatus")
HELD_COLUMNS = (
    *SOURCE_COLUMNS,
    "sourceFile",
    "sourceRow",
    "holdReason",
    "extraColumns",
    "sourceOriginal",
)
OUTPUT_NAMES = ("cleaned.csv", "needs_review.csv", "held.csv", "review.csv", "summary.json")

MISSING_MARKERS = {"", "-", "--", ".", "z", "zz"}
TAXON_MISSING_MARKERS = MISSING_MARKERS | {"sp", "sp.", "spp", "spp.", "libre", "helecho"}
FAMILY_MISSING_MARKERS = TAXON_MISSING_MARKERS
DATE_MISSING_MARKERS = MISSING_MARKERS | {"s/f", "libre", "#ref!", "z"}
MONTHS = {
    "ene": 1,
    "enero": 1,
    "jan": 1,
    "january": 1,
    "feb": 2,
    "febrero": 2,
    "february": 2,
    "mar": 3,
    "marzo": 3,
    "march": 3,
    "abr": 4,
    "abril": 4,
    "apr": 4,
    "april": 4,
    "may": 5,
    "mayo": 5,
    "jun": 6,
    "junio": 6,
    "june": 6,
    "jul": 7,
    "julio": 7,
    "july": 7,
    "ago": 8,
    "agosto": 8,
    "aug": 8,
    "august": 8,
    "set": 9,
    "sept": 9,
    "sep": 9,
    "setiembre": 9,
    "septiembre": 9,
    "september": 9,
    "oct": 10,
    "octubre": 10,
    "october": 10,
    "nov": 11,
    "noviembre": 11,
    "november": 11,
    "dic": 12,
    "diciembre": 12,
    "dec": 12,
    "december": 12,
}
COUNTRY_ALIASES = {
    "argentina": ("Argentina", "AR"),
    "bolivia": ("Bolivia", "BO"),
    "chile": ("Chile", "CL"),
    "colombia": ("Colombia", "CO"),
    "ecuador": ("Ecuador", "EC"),
    "espana": ("España", "ES"),
    "guayana francesa": ("Guayana Francesa", "GF"),
    "maine.u.s.a": ("Estados Unidos", "US"),
    "marruecos": ("Marruecos", "MA"),
    "muarruecos": ("Marruecos", "MA"),
    "panama": ("Panamá", "PA"),
    "paraguay": ("Paraguay", "PY"),
    "peru": ("Perú", "PE"),
    "rumania": ("Rumanía", "RO"),
    "suriname": ("Suriname", "SR"),
    "venezuela": ("Venezuela", "VE"),
    "grecia": ("Grecia", "GR"),
    "graecia": ("Grecia", "GR"),
}


def normalize_text(value: str) -> str:
    """Normalize Unicode and whitespace without changing source spelling."""
    return " ".join(unicodedata.normalize("NFC", value).split())


def comparison_key(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", normalize_text(value))
    return "".join(char for char in normalized if not unicodedata.combining(char)).casefold()


def clean_value(value: str, markers: Iterable[str] = MISSING_MARKERS) -> Optional[str]:
    value = normalize_text(value)
    return None if value.casefold() in markers else value


@dataclass(frozen=True)
class ReviewIssue:
    source_row: int
    catalog_number: str
    field: str
    original_value: str
    resulting_value: str
    severity: str
    reason: str


@dataclass
class CleanedRecord:
    source_row: int
    catalog_number: str
    fields: dict[str, str] = field(default_factory=dict)
    properties: dict[str, Any] = field(default_factory=dict)
    status: set[str] = field(default_factory=set)

    def set(self, key: str, value: Any) -> None:
        if value is not None:
            self.fields[key] = str(value)


@dataclass(frozen=True)
class ParsedDate:
    event_date: Optional[str]
    year: Optional[int]
    month: Optional[int]
    day: Optional[int]
    precision: str


@dataclass(frozen=True)
class Geography:
    departments: dict[str, str]
    department_ids_by_name: dict[str, set[str]]
    province_department_ids_by_name: dict[str, set[str]]
    provinces_by_name_and_department: dict[tuple[str, str], set[str]]
    districts_by_name_and_province: dict[tuple[str, str], set[str]]


class InputError(ValueError):
    """A readable input or field-map error suitable for the command line."""


def read_priorities(path: Path) -> dict[str, str]:
    """Read obligatory fields from workbook sheets with distinct layouts."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise InputError(f"No se pudo leer el mapa de campos {path}: {exc}") from exc
    required: dict[str, str] = {}
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header_row = None
            for row in rows:
                normalized = [normalize_text(str(cell or "")).casefold() for cell in row]
                if "campo en dwc" in normalized:
                    header_row = normalized
                    break
            if header_row is None or "prioridad del campo" not in header_row:
                continue
            term_index = header_row.index("campo en dwc")
            priority_index = header_row.index("prioridad del campo")
            for row in rows:
                if max(term_index, priority_index) >= len(row):
                    continue
                term_value = normalize_text(str(row[term_index] or ""))
                priority = normalize_text(str(row[priority_index] or "")).casefold()
                if priority != "obligatorio" or not term_value:
                    continue
                term = term_value.split("(", 1)[0].strip()
                if term:
                    required.setdefault(term, worksheet.title)
    finally:
        workbook.close()
    if not required:
        raise InputError(f"No se encontraron campos obligatorios en {path}.")
    return required


def load_ubigeo(data_dir: Path) -> Geography:
    """Load checked-in INEI tables and create constant-time exact lookups."""
    filenames = {
        "departments": "ubigeo_peru_2016_departamentos.json",
        "provinces": "ubigeo_peru_2016_provincias.json",
        "districts": "ubigeo_peru_2016_distritos.json",
    }
    tables: dict[str, list[dict[str, Any]]] = {}
    for key, filename in filenames.items():
        path = data_dir / filename
        if not path.is_file():
            raise InputError(f"Falta el catálogo geográfico requerido: {path}")
        try:
            tables[key] = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise InputError(f"No se pudo leer el catálogo geográfico {path}: {exc}") from exc

    departments = {item["id"]: item["name"] for item in tables["departments"]}
    department_ids_by_name: dict[str, set[str]] = defaultdict(set)
    for department_id, name in departments.items():
        department_ids_by_name[comparison_key(name)].add(department_id)

    province_department_ids_by_name: dict[str, set[str]] = defaultdict(set)
    provinces_by_name_and_department: dict[tuple[str, str], set[str]] = defaultdict(set)
    for item in tables["provinces"]:
        province_id = item["id"]
        department_id = item["department_id"]
        key = comparison_key(item["name"])
        province_department_ids_by_name[key].add(department_id)
        provinces_by_name_and_department[(key, department_id)].add(province_id)

    districts_by_name_and_province: dict[tuple[str, str], set[str]] = defaultdict(set)
    for item in tables["districts"]:
        districts_by_name_and_province[(comparison_key(item["name"]), item["province_id"])].add(
            item["department_id"]
        )

    return Geography(
        departments=departments,
        department_ids_by_name=dict(department_ids_by_name),
        province_department_ids_by_name=dict(province_department_ids_by_name),
        provinces_by_name_and_department=dict(provinces_by_name_and_department),
        districts_by_name_and_province=dict(districts_by_name_and_province),
    )


def parse_coordinate(value: str, axis: str) -> tuple[Optional[float], Optional[str]]:
    """Parse decimal degrees or clear DMS; refuse ranges/projected coordinates."""
    original = normalize_text(value)
    if not original or original.casefold() in MISSING_MARKERS:
        return None, None
    text = original.upper().replace("O", "W")
    directions = re.findall(r"[NSEW]", text)
    if len(set(directions)) > 1 or len(directions) > 1:
        return None, "La coordenada contiene direcciones incompatibles o repetidas."
    direction = directions[0] if directions else None
    numeric_text = re.sub(r"[NSEW]", " ", text).strip()
    if direction and numeric_text.endswith("."):
        numeric_text = numeric_text[:-1].rstrip()
    if direction and axis == "latitude" and direction in {"E", "W"}:
        return None, "La latitud contiene una dirección de longitud (E/W)."
    if direction and axis == "longitude" and direction in {"N", "S"}:
        return None, "La longitud contiene una dirección de latitud (N/S)."
    dms_match = re.fullmatch(
        r"(-?)\s*(\d{1,3})\s*[°º. ]\s*(\d{1,2})\s*['′’.: ]\s*"
        r"(\d{1,2}(?:[.,]\d+)?)\s*(?:[\"″]|['′’]){0,2}",
        numeric_text,
    )
    if dms_match:
        sign, degree_text, minute_text, second_text = dms_match.groups()
        degrees, minutes, seconds = (
            float(degree_text),
            float(minute_text),
            float(second_text.replace(",", ".")),
        )
        if minutes >= 60 or seconds >= 60:
            return None, "Los minutos o segundos de la coordenada exceden 59."
        value_degrees = degrees + minutes / 60 + seconds / 3600
        if sign == "-":
            value_degrees = -value_degrees
    else:
        decimal_match = re.fullmatch(r"(-?\d+(?:[.,]\d+)?)", numeric_text)
        if not decimal_match:
            return None, "Formato ambiguo, rango o coordenada proyectada; requiere revisión manual."
        number_text = decimal_match.group(1)
        if "," in number_text and "." not in number_text:
            number_text = number_text.replace(",", ".")
        value_degrees = float(number_text)

    limit = 90 if axis == "latitude" else 180
    if not math.isfinite(value_degrees) or abs(value_degrees) > limit:
        return None, f"La coordenada excede el límite de {limit} grados."
    if direction in {"S", "W"}:
        value_degrees = -abs(value_degrees)
    elif direction in {"N", "E"}:
        if value_degrees < 0:
            return None, "El signo negativo contradice la dirección N/E."
        value_degrees = abs(value_degrees)
    return value_degrees, None


def parse_elevation(value: str) -> tuple[Optional[str], Optional[str], Optional[str]]:
    original = normalize_text(value)
    if not original or original.casefold() in MISSING_MARKERS:
        return None, None, None
    match = re.fullmatch(
        r"\s*(-?\d+(?:[.,]\d+)?)\s*(?:m(?:\.s\.n\.m\.)?|msnm)?"
        r"(?:\s*[-–]\s*(-?\d+(?:[.,]\d+)?)\s*(?:m(?:\.s\.n\.m\.)?|msnm)?)?\s*",
        original,
        re.IGNORECASE,
    )
    if not match:
        return None, None, "Altitud no numérica o con formato ambiguo."
    first = float(match.group(1).replace(",", "."))
    second = float(match.group(2).replace(",", ".")) if match.group(2) else first
    low, high = sorted((first, second))
    return original, format_number(low), format_number(high)


def format_number(value: float) -> str:
    return str(int(value)) if value.is_integer() else f"{value:.8f}".rstrip("0").rstrip(".")


def parse_date(value: str) -> tuple[Optional[ParsedDate], Optional[str]]:
    original = normalize_text(value)
    if not original or original.casefold() in DATE_MISSING_MARKERS:
        return None, None
    text = original.strip().rstrip(".")
    try:
        parsed_datetime = datetime.fromisoformat(text.replace("Z", "+00:00"))
        parsed = parsed_datetime.date()
        has_time = any(
            (
                parsed_datetime.hour,
                parsed_datetime.minute,
                parsed_datetime.second,
                parsed_datetime.microsecond,
            )
        )
        event_date = parsed_datetime.isoformat() if has_time else parsed.isoformat()
        return ParsedDate(event_date, parsed.year, parsed.month, parsed.day, "day"), None
    except ValueError:
        pass
    try:
        parsed_day = date.fromisoformat(text)
        return ParsedDate(
            parsed_day.isoformat(), parsed_day.year, parsed_day.month, parsed_day.day, "day"
        ), None
    except ValueError:
        pass

    source_normalized = comparison_key(text)
    range_match = re.fullmatch(r"([a-z]+)\.?\s*[-/]\s*([a-z]+)\.?\s*(\d{4})", source_normalized)
    if range_match:
        month_a = MONTHS.get(range_match.group(1))
        month_b = MONTHS.get(range_match.group(2))
        range_year = int(range_match.group(3))
        if month_a and month_b and 1500 <= range_year <= 2100:
            start = f"{range_year:04d}-{month_a:02d}"
            end = f"{range_year:04d}-{month_b:02d}"
            return ParsedDate(f"{start}/{end}", range_year, None, None, "month-range"), None

    simplified = " ".join(re.sub(r"[.,]+", " ", source_normalized).split())
    year_match = re.fullmatch(r"(\d{4})", simplified)
    year = int(year_match.group(1)) if year_match else None
    if year is not None and not 1500 <= year <= 2100:
        return None, "El año está fuera del intervalo aceptado; no se corrigió automáticamente."

    month_match = re.fullmatch(r"([a-z]+)\s*(\d{4})", simplified)
    if month_match:
        month = MONTHS.get(month_match.group(1))
        year = int(month_match.group(2))
        if month and 1500 <= year <= 2100:
            return ParsedDate(f"{year:04d}-{month:02d}", year, month, None, "month"), None

    numeric_match = re.fullmatch(r"(\d{1,2})\.\s*(\d{1,2})\.(\d{4})", source_normalized)
    if numeric_match:
        day, month, year = map(int, numeric_match.groups())
        try:
            parsed = date(year, month, day)
        except ValueError:
            parsed = None
        if parsed and 1500 <= year <= 2100:
            return ParsedDate(parsed.isoformat(), year, month, day, "day"), None

    full_match = re.fullmatch(r"(\d{1,2})\s+([a-z]+)\s+(\d{4})", simplified)
    if full_match:
        day, month_name, year = (
            int(full_match.group(1)),
            full_match.group(2),
            int(full_match.group(3)),
        )
        month = MONTHS.get(month_name)
        try:
            parsed = date(year, month or 0, day)
        except ValueError:
            parsed = None
        if parsed:
            return ParsedDate(parsed.isoformat(), year, month, day, "day"), None

    month_day_match = re.fullmatch(r"([a-z]+)\s+(\d{1,2})\s+(\d{4})", simplified)
    if month_day_match:
        month_name, day, year = (
            month_day_match.group(1),
            int(month_day_match.group(2)),
            int(month_day_match.group(3)),
        )
        month = MONTHS.get(month_name)
        try:
            parsed = date(year, month or 0, day)
        except ValueError:
            parsed = None
        if parsed and 1500 <= year <= 2100:
            return ParsedDate(parsed.isoformat(), year, month, day, "day"), None

    if year_match and 1500 <= year <= 2100:
        return ParsedDate(str(year), year, None, None, "year"), None
    return None, "Fecha sin formato reconocido; se conserva el valor original para revisión."


def taxon_missing(value: str, *, family: bool = False) -> Optional[str]:
    markers = FAMILY_MISSING_MARKERS if family else TAXON_MISSING_MARKERS
    return clean_value(value, markers)


def parse_taxon(row: dict[str, str]) -> tuple[dict[str, str], list[tuple[str, str, str]]]:
    """Parse only clear rank and epithet patterns from the legacy columns."""
    family = taxon_missing(row["Familia"], family=True)
    raw_genus = taxon_missing(row["Género"])
    raw_species = taxon_missing(row["Especie"])
    raw_infra = taxon_missing(row["Clasificación subespecífica"])
    authorship = clean_value(row["Autor"])
    notes: list[tuple[str, str, str]] = []
    result: dict[str, str] = {}
    if family:
        result["family"] = family

    genus: Optional[str] = raw_genus
    epithet: Optional[str] = None
    infra_epithet: Optional[str] = None
    infra_rank: Optional[str] = None
    qualifier_parts: list[str] = []

    # Split an explicit qualifier from a genus, then handle clear binomials.
    if genus:
        genus_qualifier = re.fullmatch(r"(?i:cf\.?)\s+([A-Z][a-z]+)", genus)
        if genus_qualifier:
            genus = genus_qualifier.group(1)
            qualifier_parts.append("cf.")
            notes.append(
                (
                    "identificationQualifier",
                    raw_genus or "",
                    "Se separó un calificador taxonómico explícito de Género.",
                )
            )
        words = genus.split()
        if (
            len(words) == 2
            and words[0][:1].isupper()
            and words[1][:1].islower()
            and comparison_key(words[1]) not in {"sp", "spp"}
        ):
            genus, epithet = words
            notes.append(
                (
                    "Género",
                    raw_genus or "",
                    "Se separó un binomio claro registrado en la columna Género.",
                )
            )
        elif len(words) > 1:
            notes.append(
                (
                    "Género",
                    raw_genus or "",
                    "El valor de Género contiene varios términos; se conserva en dynamicProperties.",
                )
            )
            genus = None

    if raw_species:
        species = raw_species
        qualifier_match = re.match(r"^(cf\.?|aff\.?|nr\.?)\s+(.+)$", species, re.IGNORECASE)
        if qualifier_match:
            qualifier_parts.append(qualifier_match.group(1))
            species = qualifier_match.group(2).strip()
            notes.append(
                (
                    "identificationQualifier",
                    raw_species,
                    "Se separó un calificador taxonómico explícito.",
                )
            )
        # Remove a repeated genus only when one plain epithet remains.
        species_words = species.split()
        if (
            genus
            and len(species_words) == 2
            and species_words[0] == genus
            and re.fullmatch(r"[a-z][a-z-]*", species_words[1])
        ):
            species = species_words[1]
            notes.append(
                (
                    "specificEpithet",
                    raw_species,
                    "Se separó el epíteto específico de un binomio que repite Género.",
                )
            )
        rank_match = re.search(
            r"\b(subsp\.?|ssp\.?|var\.?|f\.?|forma)\s+(.+)$", species, re.IGNORECASE
        )
        if rank_match:
            marker = comparison_key(rank_match.group(1)).rstrip(".")
            infra_rank = {
                "ssp": "subspecies",
                "subsp": "subspecies",
                "var": "variety",
                "f": "form",
                "forma": "form",
            }.get(marker, marker)
            infra_epithet = rank_match.group(2).strip()
            species = species[: rank_match.start()].strip()
        if not epithet:
            words = species.split()
            if len(words) == 1:
                epithet = words[0]
            elif (
                len(words) == 2 and words[0][:1].isupper() and words[1][:1].islower() and not genus
            ):
                genus, epithet = words
                notes.append(
                    (
                        "Especie",
                        raw_species,
                        "Se separó un binomio claro registrado en la columna Especie.",
                    )
                )
            else:
                notes.append(
                    (
                        "Especie",
                        raw_species,
                        "No se pudo separar el valor de Especie sin inferir datos.",
                    )
                )
    if raw_infra:
        infra = raw_infra
        match = re.match(r"^(subsp\.?|ssp\.?|var\.?|f\.?|forma)\s+(.+)$", infra, re.IGNORECASE)
        if match:
            marker = comparison_key(match.group(1)).rstrip(".")
            infra_rank = {
                "ssp": "subspecies",
                "subsp": "subspecies",
                "var": "variety",
                "f": "form",
                "forma": "form",
            }.get(marker, marker)
            infra_epithet = match.group(2).strip()
        else:
            notes.append(
                (
                    "Clasificación subespecífica",
                    raw_infra,
                    "No se reconoció un rango infraespecífico en esta columna; "
                    "se conserva para revisión.",
                )
            )

    if genus:
        result["genus"] = genus
    if epithet:
        result["specificEpithet"] = epithet
    if infra_epithet:
        result["infraspecificEpithet"] = infra_epithet
    if infra_rank:
        result["taxonRank"] = infra_rank
    elif epithet:
        result["taxonRank"] = "species"
    elif genus:
        result["taxonRank"] = "genus"
    elif family:
        result["taxonRank"] = "family"
    if authorship:
        result["scientificNameAuthorship"] = authorship

    if qualifier_parts:
        result["identificationQualifier"] = " ".join(qualifier_parts)
    name_parts = [part for part in (genus, epithet) if part]
    if infra_epithet:
        rank_display = {"subspecies": "subsp.", "variety": "var.", "form": "f."}.get(
            infra_rank or "", ""
        )
        if rank_display:
            name_parts.extend((rank_display, infra_epithet))
        else:
            name_parts.append(infra_epithet)
    if authorship and name_parts:
        name_parts.append(authorship)
    scientific_name = " ".join(name_parts)
    if scientific_name:
        result["scientificName"] = scientific_name
    raw_identification = " ".join(
        value
        for value in (
            row["Género"],
            row["Especie"],
            row["Clasificación subespecífica"],
            row["Autor"],
        )
        if normalize_text(value)
    )
    if raw_identification:
        result["verbatimIdentification"] = normalize_text(raw_identification)
    return result, notes


def map_geography(
    row: dict[str, str], geography: Geography
) -> tuple[dict[str, str], list[tuple[str, str, str, str]]]:
    fields: dict[str, str] = {}
    notes: list[tuple[str, str, str, str]] = []
    raw_country = clean_value(row["País"])
    raw_department = clean_value(row["Departamento"])
    raw_province = clean_value(row["Provincia"])
    raw_locality = clean_value(row["Dtto./Localidad"])
    country: Optional[tuple[str, str]] = None
    explicit_country_present = raw_country is not None
    if raw_country:
        country = COUNTRY_ALIASES.get(comparison_key(raw_country))
        if country:
            fields["country"], fields["countryCode"] = country
            if raw_country != country[0]:
                notes.append(
                    ("country", raw_country, country[0], "Se normalizó un alias explícito de país.")
                )
        else:
            notes.append(
                (
                    "country",
                    raw_country,
                    "",
                    "País no reconocido por los alias conservadores; se preserva en dynamicProperties.",
                )
            )

    country_conflict = False
    country_from_department = (
        COUNTRY_ALIASES.get(comparison_key(raw_department)) if raw_department else None
    )
    if country and country_from_department and country != country_from_department:
        country_conflict = True
        notes.append(
            (
                "country",
                raw_department or "",
                country[0],
                "El país de Departamento contradice el país de País; requiere revisión.",
            )
        )
    elif country is None and country_from_department:
        country = country_from_department
        explicit_country_present = True
        fields["country"], fields["countryCode"] = country
        notes.append(
            (
                "country",
                raw_department or "",
                country[0],
                "Se interpretó un país registrado en la columna Departamento.",
            )
        )

    # The INEI hierarchy is only safe to apply to Peru records. In particular,
    # a foreign province can share a name with a Peruvian one (e.g. Santa Cruz).
    peru_geography_allowed = not country_conflict and (
        country is not None
        and country[1] == "PE"
        or country is None
        and not explicit_country_present
    )
    department_id: Optional[str] = None
    department_name: Optional[str] = None
    province_id: Optional[str] = None
    if peru_geography_allowed:
        if raw_department:
            department_key = comparison_key(raw_department)
            department_ids = geography.department_ids_by_name.get(department_key, set())
            if len(department_ids) == 1:
                department_id = next(iter(department_ids))
                department_name = geography.departments[department_id]
            elif not country_from_department:
                # A few old rows place a province in Departamento. Accept only
                # an exact province name that identifies one INEI department.
                candidates = geography.province_department_ids_by_name.get(department_key, set())
                if len(candidates) == 1:
                    department_id = next(iter(candidates))
                    department_name = geography.departments[department_id]
                    notes.append(
                        (
                            "stateProvince",
                            raw_department,
                            department_name,
                            "Se recuperó el departamento porque el valor coincide con una provincia única del catálogo INEI.",
                        )
                    )
                else:
                    notes.append(
                        (
                            "stateProvince",
                            raw_department,
                            "",
                            "Departamento no coincide exactamente con el catálogo INEI.",
                        )
                    )

        if department_name:
            fields["stateProvince"] = department_name
        if raw_province and department_id:
            province_ids = geography.provinces_by_name_and_department.get(
                (comparison_key(raw_province), department_id), set()
            )
            if len(province_ids) == 1:
                province_id = next(iter(province_ids))
                fields["county"] = raw_province
            else:
                notes.append(
                    (
                        "county",
                        raw_province,
                        "",
                        "La provincia no coincide con el departamento según el catálogo INEI.",
                    )
                )
        elif raw_province:
            candidates = geography.province_department_ids_by_name.get(
                comparison_key(raw_province), set()
            )
            if len(candidates) == 1:
                department_id = next(iter(candidates))
                department_name = geography.departments[department_id]
                province_ids = geography.provinces_by_name_and_department.get(
                    (comparison_key(raw_province), department_id), set()
                )
                if len(province_ids) == 1:
                    province_id = next(iter(province_ids))
                    fields["stateProvince"] = department_name
                    fields["county"] = raw_province
                    notes.append(
                        (
                            "stateProvince",
                            row["Departamento"],
                            department_name,
                            "Se infirió el departamento a partir de una provincia INEI única.",
                        )
                    )
            else:
                notes.append(
                    (
                        "county",
                        raw_province,
                        "",
                        "La provincia no identifica un departamento único en el catálogo INEI.",
                    )
                )
    elif raw_department or raw_province:
        notes.append(
            (
                "stateProvince",
                raw_province or raw_department or "",
                "",
                "No se aplicó el catálogo INEI porque el país está fuera de Perú o no se pudo validar; se conservó la jerarquía original.",
            )
        )

    if raw_locality:
        fields["locality"] = raw_locality
        if province_id:
            district_departments = geography.districts_by_name_and_province.get(
                (comparison_key(raw_locality), province_id), set()
            )
            if len(district_departments) == 1:
                fields["municipality"] = raw_locality
    if country is None and not explicit_country_present and department_id and province_id:
        fields["country"] = "Perú"
        fields["countryCode"] = "PE"
        notes.append(
            (
                "country",
                "",
                "Perú",
                "Inferido por coincidencia exacta de departamento y provincia en el catálogo INEI.",
            )
        )
    return fields, notes


def source_dict(values: Sequence[str]) -> dict[str, str]:
    return dict(zip(SOURCE_COLUMNS, values))


def clean_record(
    row: dict[str, str], source_row: int, geography: Geography
) -> tuple[CleanedRecord, list[ReviewIssue]]:
    catalog = normalize_text(row["Código USM"])
    record = CleanedRecord(source_row, catalog)
    issues: list[ReviewIssue] = []
    record.set("catalogNumber", catalog)

    taxon_fields, taxon_notes = parse_taxon(row)
    record.fields.update(taxon_fields)
    for field_name, original, reason in taxon_notes:
        severity = (
            "info"
            if field_name == "identificationQualifier" or reason.startswith("Se separó")
            else "review"
        )
        issues.append(
            ReviewIssue(
                source_row,
                catalog,
                field_name,
                original,
                taxon_fields.get(field_name, ""),
                severity,
                reason,
            )
        )

    vernacular = clean_value(row["Nombre vernáculo"])
    if vernacular:
        record.set("vernacularName", vernacular)
    description = clean_value(row["Descripción"])
    if description:
        record.set("occurrenceRemarks", description)
    record.properties["sourceData"] = row

    specimen_type = clean_value(row["Tipo"])
    if specimen_type:
        normalized_type = comparison_key(specimen_type)
        type_aliases = {
            "holotipo": "holotype",
            "holotype": "holotype",
            "isotipo": "isotype",
            "isotype": "isotype",
            "paratipo": "paratype",
            "paratype": "paratype",
            "lectotipo": "lectotype",
            "lectotype": "lectotype",
            "neotipo": "neotype",
            "neotype": "neotype",
        }
        if normalized_type in type_aliases:
            record.set("typeStatus", type_aliases[normalized_type])
        else:
            issues.append(
                ReviewIssue(
                    source_row,
                    catalog,
                    "typeStatus",
                    specimen_type,
                    "",
                    "review",
                    "Tipo no coincide con un estatus tipológico reconocido; se conserva en sourceData.",
                )
            )

    geography_fields, geography_notes = map_geography(row, geography)
    record.fields.update(geography_fields)
    for field_name, original, resulting, reason in geography_notes:
        reason_key = reason.casefold()
        severity = (
            "info"
            if any(
                token in reason_key for token in ("normalizó", "interpretó", "infer", "recuperó")
            )
            else "review"
        )
        issues.append(
            ReviewIssue(source_row, catalog, field_name, original, resulting, severity, reason)
        )
    if row["Dtto./Localidad"]:
        record.set("verbatimLocality", normalize_text(row["Dtto./Localidad"]))

    lat, lat_problem = parse_coordinate(row["Latitud / Northing"], "latitude")
    lon, lon_problem = parse_coordinate(row["Longitud / Easting"], "longitude")
    if lat_problem:
        issues.append(
            ReviewIssue(
                source_row,
                catalog,
                "decimalLatitude",
                row["Latitud / Northing"],
                "",
                "review",
                lat_problem,
            )
        )
    if lon_problem:
        issues.append(
            ReviewIssue(
                source_row,
                catalog,
                "decimalLongitude",
                row["Longitud / Easting"],
                "",
                "review",
                lon_problem,
            )
        )
    if (lat is None) != (lon is None):
        issues.append(
            ReviewIssue(
                source_row,
                catalog,
                "coordinates",
                f"{row['Latitud / Northing']} | {row['Longitud / Easting']}",
                "",
                "review",
                "Se requiere un par completo de latitud y longitud; se dejaron ambas coordenadas vacías.",
            )
        )
        lat = lon = None
    if lat is not None and lon is not None:
        record.set("decimalLatitude", format_number(lat))
        record.set("decimalLongitude", format_number(lon))

    elevation, elevation_min, elevation_max = parse_elevation(row["Altitud"])
    if elevation:
        record.set("verbatimElevation", elevation)
        record.set("minimumElevationInMeters", elevation_min)
        record.set("maximumElevationInMeters", elevation_max)
        if "-" in elevation or "–" in elevation:
            issues.append(
                ReviewIssue(
                    source_row,
                    catalog,
                    "minimumElevationInMeters",
                    row["Altitud"],
                    f"{elevation_min}/{elevation_max}",
                    "info",
                    "Se separó un intervalo de altitud explícito.",
                )
            )
    elif elevation_problem := elevation_min:
        issues.append(
            ReviewIssue(
                source_row,
                catalog,
                "verbatimElevation",
                row["Altitud"],
                "",
                "review",
                elevation_problem,
            )
        )

    collector = clean_value(row["Colector"])
    if collector:
        record.set("recordedBy", collector)
    collector_number = clean_value(row["N° Colector"])
    if collector_number:
        record.set("recordNumber", collector_number)

    raw_event_date = row["Fecha colecta"]
    record.set("verbatimEventDate", clean_value(raw_event_date))
    parsed_date, date_problem = parse_date(raw_event_date)
    if parsed_date:
        record.set("eventDate", parsed_date.event_date)
        record.set("year", parsed_date.year)
        record.set("month", parsed_date.month)
        record.set("day", parsed_date.day)
        if normalize_text(raw_event_date) != parsed_date.event_date:
            issues.append(
                ReviewIssue(
                    source_row,
                    catalog,
                    "eventDate",
                    raw_event_date,
                    parsed_date.event_date or "",
                    "info",
                    "Se normalizó la fecha conservando la precisión disponible.",
                )
            )
    elif date_problem:
        issues.append(
            ReviewIssue(
                source_row, catalog, "eventDate", raw_event_date, "", "review", date_problem
            )
        )

    identified_by = clean_value(row["Determinado por"])
    if identified_by:
        record.set("identifiedBy", identified_by)
    date_identified_raw = row["Fecha determinacion"]
    parsed_identified, identified_problem = parse_date(date_identified_raw)
    if parsed_identified:
        record.set("dateIdentified", parsed_identified.event_date)
        if normalize_text(date_identified_raw) != parsed_identified.event_date:
            issues.append(
                ReviewIssue(
                    source_row,
                    catalog,
                    "dateIdentified",
                    date_identified_raw,
                    parsed_identified.event_date or "",
                    "info",
                    "Se normalizó la fecha de determinación conservando su precisión.",
                )
            )
    elif identified_problem:
        issues.append(
            ReviewIssue(
                source_row,
                catalog,
                "dateIdentified",
                date_identified_raw,
                "",
                "review",
                identified_problem,
            )
        )

    properties = record.properties
    if clean_value(row["Otros"]):
        properties["sourceOther"] = normalize_text(row["Otros"])
    for source_column in (
        "Registrador",
        "Fecha digitalización",
        "Distribuidos en .. (Herbarios)",
        "Fecha de revisión",
    ):
        source_value = clean_value(row[source_column])
        if source_value and not (
            source_column == "Registrador" and comparison_key(source_value) == "registrador"
        ):
            properties[source_column] = source_value
    if properties:
        record.set(
            "dynamicProperties", json.dumps(properties, ensure_ascii=False, separators=(",", ":"))
        )

    if any(issue.severity == "review" for issue in issues):
        record.status.add("review")
    return record, issues


def hold_reason(row: dict[str, str]) -> Optional[str]:
    catalog = normalize_text(row.get("Código USM", ""))
    if comparison_key(catalog) == "libre":
        return "Código USM marcado como LIBRE; no representa un espécimen para importar."
    if comparison_key(row.get("Familia", "")) == "libre":
        return "Fila de reserva marcada como LIBRE en Familia; no representa un espécimen para importar."
    if not re.fullmatch(r"\d{1,100}", catalog):
        return "Código USM vacío o inválido; se esperaba de 1 a 100 dígitos."
    return None


@dataclass(frozen=True)
class SourceLayout:
    headers: tuple[str, ...]
    targets: tuple[Optional[str], ...]
    adaptations: tuple[str, ...]

    def map_row(self, values: Sequence[str]) -> dict[str, str]:
        row = dict.fromkeys(SOURCE_COLUMNS, "")
        for target, value in zip(self.targets, values):
            if target is not None:
                row[target] = value
        return row

    def original(self, values: Sequence[str]) -> dict[str, Any]:
        return {"headers": list(self.headers), "values": list(values)}

    def extra_cells(self, values: Sequence[str]) -> list[dict[str, Any]]:
        return [
            {"column": header, "index": index, "value": values[index]}
            for index, (header, target) in enumerate(zip(self.headers, self.targets))
            if target is None and index < len(values)
        ]


def read_layout(path: Path) -> SourceLayout:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            headers = tuple(next(csv.reader(handle, strict=True)))
    except StopIteration as exc:
        raise InputError(f"El CSV está vacío: {path}") from exc
    except (OSError, UnicodeError, csv.Error) as exc:
        raise InputError(f"No se pudo leer el CSV {path}: {exc}") from exc
    names = [normalize_text(header).casefold() for header in headers]
    if len(set(names)) != len(names):
        raise InputError(f"Encabezados repetidos después de normalizar en {path.name}.")
    canonical = {name.casefold(): name for name in SOURCE_COLUMNS}
    targets = [canonical.get(name) for name in names]
    adaptations = [
        f"{header!r} -> {target!r}"
        for header, target in zip(headers, targets)
        if target is not None and header != target
    ]
    # This unnamed column is a verified accession column, not an arbitrary index.
    expected_8 = tuple(name for name in SOURCE_COLUMNS if name != "Clasificación subespecífica")
    if (
        path.name == "8_70.csv"
        and names
        and names[0] == "unnamed: 0"
        and tuple(targets[1:]) == expected_8[1:]
    ):
        targets[0] = "Código USM"
        adaptations.append("Unnamed: 0 -> Código USM (perfil 8_70.csv)")
    missing = set(SOURCE_COLUMNS) - set(targets)
    required_missing = missing - {"Autor", "Clasificación subespecífica"}
    if required_missing:
        raise InputError(
            f"Encabezados requeridos faltantes en {path.name}: "
            + ", ".join(sorted(required_missing))
        )
    adaptations.extend(f"Columna ausente, valor vacío: {name}" for name in sorted(missing))
    if tuple(targets) != SOURCE_COLUMNS and not missing and all(targets):
        adaptations.append("Columnas reordenadas por nombre")
    return SourceLayout(headers, tuple(targets), tuple(adaptations))


def iter_source_rows(path: Path, layout: SourceLayout) -> Iterable[tuple[int, list[str]]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, strict=True)
            header = tuple(next(reader))
            if header != layout.headers:
                raise InputError(f"Los encabezados cambiaron durante la lectura: {path}")
            for values in reader:
                # Keep the established physical end-line reference for multiline CSV records.
                yield reader.line_num, values
    except (OSError, UnicodeError, csv.Error, StopIteration) as exc:
        raise InputError(f"No se pudo leer el CSV {path}: {exc}") from exc


def read_source(path: Path) -> tuple[list[tuple[int, list[str]]], list[tuple[int, list[str]]]]:
    """Read rows classified against their actual source header width."""
    layout = read_layout(path)
    valid: list[tuple[int, list[str]]] = []
    malformed: list[tuple[int, list[str]]] = []
    for source_row, values in iter_source_rows(path, layout):
        (valid if len(values) == len(layout.headers) else malformed).append((source_row, values))
    return valid, malformed


def apply_source_context(
    record: CleanedRecord,
    issues: list[ReviewIssue],
    row: dict[str, str],
    path: Path,
    layout: SourceLayout,
    values: Sequence[str],
) -> None:
    record.properties["sourceOriginal"] = layout.original(values)
    extras = layout.extra_cells(values)
    meaningful = []
    for cell in extras:
        value = clean_value(cell["value"])
        if value is None:
            continue
        if (
            path.name == "35_340.csv"
            and cell["column"] == "Unnamed: 26"
            and value == normalize_text(row["Fecha de revisión"])
        ):
            issues.append(
                ReviewIssue(
                    record.source_row,
                    record.catalog_number,
                    "sourceExtraColumns",
                    json.dumps([cell], ensure_ascii=False),
                    "",
                    "info",
                    "Fecha adicional idéntica a Fecha de revisión; se conservó el original.",
                )
            )
            continue
        if (
            path.name == "27_260.csv"
            and cell["column"] == "Unnamed: 29"
            and value.casefold() == "isotipo"
            and clean_value(row["Tipo"]) is None
        ):
            record.set("typeStatus", "isotype")
            issues.append(
                ReviewIssue(
                    record.source_row,
                    record.catalog_number,
                    "typeStatus",
                    cell["value"],
                    "isotype",
                    "info",
                    "Se recuperó ISOTIPO de Unnamed: 29 (perfil 27_260.csv).",
                )
            )
        meaningful.append(cell)
    if extras:
        record.properties["sourceExtraColumns"] = extras
    if meaningful:
        issues.append(
            ReviewIssue(
                record.source_row,
                record.catalog_number,
                "sourceExtraColumns",
                json.dumps(meaningful, ensure_ascii=False),
                "",
                "review",
                "Columnas adicionales con contenido; revisar su significado y posibles valores desplazados.",
            )
        )
        record.status.add("review")
    record.set(
        "dynamicProperties",
        json.dumps(record.properties, ensure_ascii=False, separators=(",", ":")),
    )


def write_csv(path: Path, headers: Sequence[str], rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def select_batch_sources(input_dir: Path) -> tuple[list[Path], list[str]]:
    if not input_dir.is_dir():
        raise InputError(f"No existe el directorio de entrada: {input_dir}")
    selected: dict[int, list[Path]] = defaultdict(list)
    excluded = []
    for path in sorted(input_dir.glob("*.csv")):
        match = re.match(r"^(\d+)_", path.name)
        if match and 1 <= int(match.group(1)) <= 37:
            selected[int(match.group(1))].append(path)
        else:
            excluded.append(path.name)
    invalid = [str(number) for number in range(1, 38) if len(selected[number]) != 1]
    if invalid:
        raise InputError(
            "Se requiere exactamente un CSV por prefijo 1–37; faltante o ambiguo: "
            + ", ".join(invalid)
        )
    return [selected[number][0] for number in range(1, 38)], excluded


def run_cleaner(
    input_path: Path,
    fields_path: Path,
    ubigeo_dir: Path,
    output_dir: Path,
    overwrite: bool = False,
) -> dict[str, Any]:
    return run_sources([input_path], fields_path, ubigeo_dir, output_dir, overwrite)


def run_batch_cleaner(
    input_dir: Path,
    fields_path: Path,
    ubigeo_dir: Path,
    output_dir: Path,
    overwrite: bool = False,
) -> dict[str, Any]:
    paths, excluded = select_batch_sources(input_dir)
    return run_sources(paths, fields_path, ubigeo_dir, output_dir, overwrite, excluded)


def run_sources(
    input_paths: Sequence[Path],
    fields_path: Path,
    ubigeo_dir: Path,
    output_dir: Path,
    overwrite: bool = False,
    excluded: Sequence[str] = (),
) -> dict[str, Any]:
    for path in (*input_paths, fields_path):
        if not path.is_file():
            raise InputError(f"No existe el archivo requerido: {path}")
    existing = [output_dir / name for name in OUTPUT_NAMES if (output_dir / name).exists()]
    if existing and not overwrite:
        raise InputError(
            "Ya existen archivos de salida; use --overwrite para reemplazarlos: "
            + ", ".join(map(str, existing))
        )
    if len({path.name for path in input_paths}) != len(input_paths):
        raise InputError("Los archivos de entrada deben tener nombres distintos.")
    layouts = {path: read_layout(path) for path in input_paths}
    priorities = read_priorities(fields_path)
    geography = load_ubigeo(ubigeo_dir)
    duplicate_groups: dict[str, list[tuple[str, int]]] = defaultdict(list)
    # Preflight all CSVs and index accession numbers before any report is published.
    for path in input_paths:
        layout = layouts[path]
        for source_row, values in iter_source_rows(path, layout):
            if len(values) != len(layout.headers) or not any(value.strip() for value in values):
                continue
            catalog = normalize_text(layout.map_row(values)["Código USM"])
            if re.fullmatch(r"\d{1,100}", catalog):
                duplicate_groups[catalog].append((path.name, source_row))
    duplicates = {catalog: refs for catalog, refs in duplicate_groups.items() if len(refs) > 1}
    del duplicate_groups
    counts: Counter[str] = Counter()
    severity_counts: Counter[str] = Counter()
    field_counts: Counter[str] = Counter()
    hold_reasons: Counter[str] = Counter()
    required_missing: Counter[str] = Counter()
    per_file = []
    output_terms = set(OUTPUT_DWC_FIELDS)
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="usm-clean-", dir=output_dir) as temporary_directory:
        temporary = Path(temporary_directory)
        with ExitStack() as stack:
            writers = {}
            for name, columns in (
                ("cleaned.csv", (*OUTPUT_DWC_FIELDS, *AUDIT_COLUMNS)),
                ("needs_review.csv", (*OUTPUT_DWC_FIELDS, *AUDIT_COLUMNS)),
                ("held.csv", HELD_COLUMNS),
                ("review.csv", REVIEW_COLUMNS),
            ):
                handle = stack.enter_context(
                    (temporary / name).open("w", encoding="utf-8-sig", newline="")
                )
                writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
                writer.writeheader()
                writers[name] = writer
            for path in input_paths:
                layout = layouts[path]
                file_counts: Counter[str] = Counter()
                file_duplicate_catalogs = set()
                for source_row, values in iter_source_rows(path, layout):
                    file_counts["sourceRows"] += 1
                    if not any(value.strip() for value in values):
                        file_counts["skippedBlankRows"] += 1
                        continue
                    row = layout.map_row(values)
                    catalog = normalize_text(row["Código USM"])
                    reason = None
                    if len(values) != len(layout.headers):
                        reason = (
                            f"Fila con {len(values)} columnas; se esperaban {len(layout.headers)}."
                        )
                        file_counts["malformedWidthRows"] += 1
                    else:
                        reason = hold_reason(row)
                        if catalog in duplicates:
                            reason = (
                                f"Código USM duplicado ({catalog}); "
                                "se retuvieron todas las filas del grupo."
                            )
                            file_duplicate_catalogs.add(catalog)
                            file_counts["duplicateRows"] += 1
                    if reason:
                        writers["held.csv"].writerow(
                            {
                                **row,
                                "sourceFile": path.name,
                                "sourceRow": source_row,
                                "holdReason": reason,
                                "extraColumns": json.dumps(
                                    layout.extra_cells(values)
                                    + [
                                        {"column": None, "index": index, "value": value}
                                        for index, value in enumerate(values)
                                        if index >= len(layout.headers)
                                    ],
                                    ensure_ascii=False,
                                ),
                                "sourceOriginal": json.dumps(
                                    layout.original(values), ensure_ascii=False
                                ),
                            }
                        )
                        file_counts["heldRows"] += 1
                        hold_reasons[reason] += 1
                        continue
                    record, issues = clean_record(row, source_row, geography)
                    apply_source_context(record, issues, row, path, layout, values)
                    needs_review = "review" in record.status
                    name = "needs_review.csv" if needs_review else "cleaned.csv"
                    writers[name].writerow(
                        {
                            **{term: record.fields.get(term, "") for term in OUTPUT_DWC_FIELDS},
                            "sourceFile": path.name,
                            "sourceRow": source_row,
                            "cleaningStatus": "review" if needs_review else "cleaned",
                        }
                    )
                    file_counts["reviewRows" if needs_review else "cleanedRows"] += 1
                    if not needs_review:
                        for term in priorities:
                            if term not in output_terms or not record.fields.get(term):
                                required_missing[term] += 1
                    for issue in issues:
                        writers["review.csv"].writerow(
                            {
                                "sourceFile": path.name,
                                "sourceRow": issue.source_row,
                                "catalogNumber": issue.catalog_number,
                                "field": issue.field,
                                "originalValue": issue.original_value,
                                "resultingValue": issue.resulting_value,
                                "severity": issue.severity,
                                "reason": issue.reason,
                            }
                        )
                        severity_counts[issue.severity] += 1
                        field_counts[issue.field] += 1
                        file_counts["reviewIssues"] += 1
                counts.update(file_counts)
                per_file.append(
                    {
                        "sourceFile": path.name,
                        **{
                            key: file_counts[key]
                            for key in (
                                "sourceRows",
                                "cleanedRows",
                                "reviewRows",
                                "heldRows",
                                "skippedBlankRows",
                                "malformedWidthRows",
                                "duplicateRows",
                                "reviewIssues",
                            )
                        },
                        "duplicateGroupCount": len(file_duplicate_catalogs),
                        "headerAdaptations": list(layout.adaptations),
                        "extraHeaders": [
                            header
                            for header, target in zip(layout.headers, layout.targets)
                            if target is None
                        ],
                    }
                )
        if counts["sourceRows"] != sum(
            counts[key] for key in ("cleanedRows", "reviewRows", "heldRows", "skippedBlankRows")
        ):
            raise InputError("Los totales de filas no coinciden; no se publicaron los informes.")
        summary = {
            **{
                key: counts[key]
                for key in (
                    "sourceRows",
                    "cleanedRows",
                    "reviewRows",
                    "heldRows",
                    "skippedBlankRows",
                    "malformedWidthRows",
                    "reviewIssues",
                    "duplicateRows",
                )
            },
            "sourceFiles": [str(path) for path in input_paths],
            "excludedFiles": list(excluded),
            "perFile": per_file,
            "recordsNeedingReview": counts["reviewRows"],
            "recordsWithoutReviewFlags": counts["cleanedRows"],
            "duplicateGroupCount": len(duplicates),
            "crossFileDuplicateGroupCount": sum(
                len({file for file, _ in refs}) > 1 for refs in duplicates.values()
            ),
            "issuesBySeverity": dict(severity_counts),
            "issuesByField": dict(field_counts),
            "holdReasons": dict(hold_reasons),
            "requiredFields": {
                "rowScope": "cleaned.csv",
                "byTermMissingRows": {term: required_missing[term] for term in sorted(priorities)},
                "byTermSourceSheet": priorities,
                "termsWithoutOutputMapping": sorted(set(priorities) - output_terms),
            },
            "rules": [
                "Se conservan sourceData canónico y sourceOriginal exacto en dynamicProperties.",
                "No se hizo correspondencia de taxones contra la base de datos.",
                "Solo se infirió Perú con coincidencia exacta de departamento y provincia en INEI.",
                "No se adivinaron coordenadas, rangos, formatos ambiguos ni desplazamientos de celdas.",
                "Las filas con señales de revisión se separan en needs_review.csv.",
                "Filas incompletas sin señales de revisión permanecen en cleaned.csv.",
                "Códigos inválidos, LIBRE, duplicados y anchos malformados se retienen en held.csv.",
                "Filas completamente vacías se omiten y se cuentan por separado.",
                "Todos los miembros de duplicados se retienen, también entre archivos del lote.",
            ],
            "outputs": {name: str(output_dir / name) for name in OUTPUT_NAMES},
        }
        if len(input_paths) == 1:
            summary["sourceFile"] = str(input_paths[0])
        (temporary / "summary.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        for name in OUTPUT_NAMES:
            (temporary / name).replace(output_dir / name)
    return summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    inputs = parser.add_mutually_exclusive_group()
    inputs.add_argument("--input", type=Path, help="CSV export to clean (default: 32_310.csv)")
    inputs.add_argument("--input-dir", type=Path, help="Directory of specimen CSVs 1–37")
    parser.add_argument(
        "--fields", type=Path, default=DEFAULT_FIELDS, help="Campos DwC.xlsx priority map"
    )
    parser.add_argument(
        "--ubigeo-dir",
        type=Path,
        default=DEFAULT_UBIGEO,
        help="Directory containing checked-in INEI JSON files",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        help="Directory for reports (default: data/cleaned/32_310 or data/cleaned/usm_all)",
    )
    parser.add_argument("--overwrite", action="store_true", help="Replace existing report files")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    output_dir = args.output_dir or (
        DEFAULT_BATCH_OUTPUT_DIR if args.input_dir else DEFAULT_OUTPUT_DIR
    )
    try:
        runner = run_batch_cleaner if args.input_dir else run_cleaner
        summary = runner(
            args.input_dir or args.input or DEFAULT_INPUT,
            args.fields,
            args.ubigeo_dir,
            output_dir,
            args.overwrite,
        )
    except InputError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    except OSError as exc:
        print(f"Error de lectura o escritura: {exc}", file=sys.stderr)
        return 2
    print(
        json.dumps(
            {
                key: summary[key]
                for key in (
                    "sourceRows",
                    "cleanedRows",
                    "reviewRows",
                    "heldRows",
                    "skippedBlankRows",
                    "recordsNeedingReview",
                    "reviewIssues",
                )
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    print(f"Informes guardados en: {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
